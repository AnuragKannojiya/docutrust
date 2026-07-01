import os
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether, ListFlowable, ListItem
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

DOCS_DIR = Path(__file__).parent.parent / "docs"
DOCS_DIR.mkdir(exist_ok=True)

DOCX_PATH = DOCS_DIR / "docutrust_documentation.docx"
PDF_PATH = DOCS_DIR / "docutrust_documentation.pdf"

# Content definition
SECTIONS = [
    {
        "type": "title",
        "text": "DocuTrust: Enterprise-Grade Corrective RAG (CRAG) System"
    },
    {
        "type": "subtitle",
        "text": "Technical Design and User Manual"
    },
    {
        "type": "heading1",
        "text": "1. Executive Summary"
    },
    {
        "type": "paragraph",
        "text": "DocuTrust is a self-correcting enterprise Retrieval-Augmented Generation (RAG) system. Typical RAG pipelines suffer from hallucinations, database search noise, and incomplete internal corpora. DocuTrust addresses these weaknesses using a Corrective RAG (CRAG) architecture that incorporates local vector search, Cross-Encoder relevance grading, query re-writing, allowlisted web fallback searching, and strict post-generation citation validation."
    },
    {
        "type": "heading1",
        "text": "2. System Architecture & Components"
    },
    {
        "type": "paragraph",
        "text": "The DocuTrust system consists of five tightly integrated layers, executing in a linear-with-conditional-branch pipeline managed via FastAPI and LangGraph:"
    },
    {
        "type": "heading2",
        "text": "2.1 Ingestion Layer"
    },
    {
        "type": "paragraph",
        "text": "During ingestion, uploaded documents (PDFs, Word documents, or plain text) are parsed and split into overlapping semantic paragraphs. Each chunk is assigned metadata including its parent document ID, section path, and source details. The raw text and metadata are saved to MongoDB, while the text is embedded using OpenAI's 'text-embedding-3-small' model and stored in a local FAISS index."
    },
    {
        "type": "heading2",
        "text": "2.2 Retrieval Layer"
    },
    {
        "type": "paragraph",
        "text": "For every incoming query, the question is vectorized using the same OpenAI embeddings model. A similarity search is performed against the FAISS index to retrieve the top K nearest chunks, which are then hydrated with their text content and metadata from MongoDB."
    },
    {
        "type": "heading2",
        "text": "2.3 Cross-Encoder relevance grading"
    },
    {
        "type": "paragraph",
        "text": "Rather than relying solely on vector similarity, DocuTrust routes all retrieved chunks through a local Cross-Encoder model ('cross-encoder/ms-marco-MiniLM-L-6-v2'). The Cross-Encoder performs self-attention over the query and text together, assigning a normalized score from 0.0 to 1.0. Chunks scoring below the grader threshold (default 0.55) are classified as irrelevant and discarded."
    },
    {
        "type": "heading2",
        "text": "2.4 Corrective Web Fallback"
    },
    {
        "type": "paragraph",
        "text": "If no chunks pass the Cross-Encoder grading threshold, the internal corpus is deemed insufficient. The query is automatically sent to the LLM (gpt-4o-mini) to be rewritten into optimized search engine terms. A search is executed on the open web using DuckDuckGo, restricted to high-quality allowed domains (such as sec.gov, nist.gov, and europa.eu). These web search results are formatted into chunks and used as context."
    },
    {
        "type": "heading2",
        "text": "2.5 Generation & Gating Layer"
    },
    {
        "type": "paragraph",
        "text": "The LLM generates a response using the compiled context and is instructed to cite references as '[chunk_id]'. The backend parses the LLM output; if the answer lacks valid citation markers, contains hallucinated citations, or prefixes with 'REFUSE', the system blocks the response and returns a formal refusal, ensuring zero hallucinations."
    },
    {
        "type": "heading1",
        "text": "3. Setup and Configuration"
    },
    {
        "type": "paragraph",
        "text": "Configure the application via a '.env' file located in the project root. Below is a sample configuration template:"
    },
    {
        "type": "code",
        "text": "OPENAI_API_KEY=your-openai-api-key\nMONGO_URI=mongodb://localhost:27017\nMONGO_DB=docutrust\nVECTOR_BACKEND=faiss\nFAISS_INDEX_PATH=./data/faiss\nOPENAI_CHAT_MODEL=gpt-4o-mini\nOPENAI_EMBED_MODEL=text-embedding-3-small\nGRADER_MODEL=cross-encoder/ms-marco-MiniLM-L-6-v2\nGRADER_THRESHOLD=0.55\nRETRIEVAL_K=8\nWEB_DOMAINS_ALLOWLIST=sec.gov,europa.eu,nist.gov,iso.org\nWEB_MAX_RESULTS=5"
    },
    {
        "type": "heading1",
        "text": "4. Local Execution & Ingestion Commands"
    },
    {
        "type": "paragraph",
        "text": "Follow these command-line instructions to run and populate the project locally:"
    },
    {
        "type": "bullet",
        "text": "Start the Backend API Server:"
    },
    {
        "type": "code",
        "text": ".venv/bin/uvicorn app.main:app --reload --port 8000"
    },
    {
        "type": "bullet",
        "text": "Seed the database and ingest demo files:"
    },
    {
        "type": "code",
        "text": ".venv/bin/python -m scripts.seed_corpus"
    },
    {
        "type": "bullet",
        "text": "Start the Streamlit Frontend application:"
    },
    {
        "type": "code",
        "text": ".venv/bin/streamlit run ui/streamlit_app.py"
    },
    {
        "type": "bullet",
        "text": "Execute the testing suite:"
    },
    {
        "type": "code",
        "text": ".venv/bin/pytest"
    }
]

def generate_docx():
    print(f"Generating DOCX at {DOCX_PATH}...")
    doc = Document()
    
    # Document styling settings
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Arial'
    style_normal.font.size = Pt(11)
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(6)
    
    for section in SECTIONS:
        t = section["type"]
        val = section["text"]
        
        if t == "title":
            p = doc.add_paragraph()
            p.alignment = 1 # Center
            run = p.add_run(val)
            run.font.name = 'Arial'
            run.font.size = Pt(22)
            run.bold = True
            run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D) # Navy
            p.paragraph_format.space_after = Pt(12)
            
        elif t == "subtitle":
            p = doc.add_paragraph()
            p.alignment = 1
            run = p.add_run(val)
            run.font.name = 'Arial'
            run.font.size = Pt(14)
            run.italic = True
            run.font.color.rgb = RGBColor(0x5C, 0x76, 0x8D) # Slate Gray
            p.paragraph_format.space_after = Pt(24)
            
        elif t == "heading1":
            p = doc.add_heading(level=1)
            run = p.runs[0] if p.runs else p.add_run(val)
            run.font.name = 'Arial'
            run.font.size = Pt(16)
            run.bold = True
            run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            
        elif t == "heading2":
            p = doc.add_heading(level=2)
            run = p.runs[0] if p.runs else p.add_run(val)
            run.font.name = 'Arial'
            run.font.size = Pt(13)
            run.bold = True
            run.font.color.rgb = RGBColor(0x2E, 0x5B, 0x88)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            
        elif t == "paragraph":
            p = doc.add_paragraph(val)
            
        elif t == "bullet":
            p = doc.add_paragraph(style='List Bullet')
            run = p.add_run(val)
            p.paragraph_format.space_after = Pt(3)
            
        elif t == "code":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.right_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            
            # Shading & border style using a simple grey table look
            # For simplicity, we write it in a styled single-run block with courier font.
            run = p.add_run(val)
            run.font.name = 'Courier New'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
            
    doc.save(str(DOCX_PATH))
    print("DOCX generation complete.")

def generate_pdf():
    print(f"Generating PDF at {PDF_PATH}...")
    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=letter,
        rightMargin=54, leftMargin=54,
        topMargin=54, bottomMargin=54
    )
    
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#1B365D'),
        alignment=1, # Center
        spaceAfter=12
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#5C768D'),
        alignment=1,
        spaceAfter=24
    )
    
    h1_style = ParagraphStyle(
        'DocH1',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=15,
        leading=18,
        textColor=colors.HexColor('#1B365D'),
        spaceBefore=16,
        spaceAfter=8,
        keepWithNext=True
    )
    
    h2_style = ParagraphStyle(
        'DocH2',
        parent=styles['Heading3'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=15,
        textColor=colors.HexColor('#2E5B88'),
        spaceBefore=12,
        spaceAfter=6,
        keepWithNext=True
    )
    
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['BodyText'],
        fontName='Helvetica',
        fontSize=10,
        leading=14.5,
        textColor=colors.HexColor('#333333'),
        spaceAfter=8
    )
    
    bullet_style = ParagraphStyle(
        'DocBullet',
        parent=body_style,
        leftIndent=20,
        firstLineIndent=-10,
        spaceAfter=4
    )
    
    code_style = ParagraphStyle(
        'DocCode',
        parent=styles['Code'],
        fontName='Courier',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor('#111111'),
        backColor=colors.HexColor('#F5F5F5'),
        borderColor=colors.HexColor('#E0E0E0'),
        borderWidth=0.5,
        borderPadding=8,
        spaceBefore=6,
        spaceAfter=6,
        leftIndent=15,
        rightIndent=15
    )

    story = []
    
    for section in SECTIONS:
        t = section["type"]
        val = section["text"]
        
        if t == "title":
            story.append(Paragraph(val, title_style))
        elif t == "subtitle":
            story.append(Paragraph(val, subtitle_style))
        elif t == "heading1":
            story.append(Paragraph(val, h1_style))
        elif t == "heading2":
            story.append(Paragraph(val, h2_style))
        elif t == "paragraph":
            story.append(Paragraph(val, body_style))
        elif t == "bullet":
            story.append(Paragraph(f"&bull; {val}", bullet_style))
        elif t == "code":
            # Format raw code text to escape html characters for reportlab
            escaped_val = val.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>")
            story.append(Paragraph(escaped_val, code_style))
            
    doc.build(story)
    print("PDF generation complete.")

if __name__ == "__main__":
    generate_docx()
    generate_pdf()
